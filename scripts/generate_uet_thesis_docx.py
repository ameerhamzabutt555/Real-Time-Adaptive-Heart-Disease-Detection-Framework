#!/usr/bin/env python3
"""Generate a UET/ORIC-aligned thesis draft DOCX from implemented project content."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def add_centered(doc: Document, text: str, bold: bool = False) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = para.add_run(text)
    run.bold = bold


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_section(doc: Document, heading: str, paragraphs: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)


def build_document() -> Document:
    doc = Document()

    # Title page
    add_centered(
        doc,
        "A REAL-TIME ADAPTIVE FRAMEWORK FOR DETECTION OF HEART DISEASES",
        bold=True,
    )
    doc.add_paragraph("")
    add_centered(doc, "by")
    add_centered(doc, "[FULL NAME OF THE SCHOLAR]")
    add_centered(doc, "[REGISTRATION NUMBER]")
    doc.add_paragraph("")
    add_centered(doc, "Research Supervisor:")
    add_centered(doc, "[NAME OF RESEARCH SUPERVISOR]")
    doc.add_paragraph("")
    add_centered(doc, "[YEAR]")
    doc.add_paragraph("")
    add_centered(doc, "Department of [DEPARTMENT NAME]")
    add_centered(doc, "University of Engineering and Technology, Lahore")
    add_page_break(doc)

    # Approval page
    add_centered(
        doc,
        "A REAL-TIME ADAPTIVE FRAMEWORK FOR DETECTION OF HEART DISEASES",
        bold=True,
    )
    doc.add_paragraph("")
    add_centered(doc, "by")
    add_centered(doc, "[FULL NAME OF THE SCHOLAR]")
    doc.add_paragraph("")
    add_centered(doc, "A THESIS")
    add_centered(
        doc,
        "presented to the University of Engineering and Technology, Lahore in partial fulfillment",
    )
    add_centered(doc, "of the requirements for the degree of Master of Philosophy in [SUBJECT]")
    doc.add_paragraph("")
    doc.add_paragraph("APPROVED BY:")
    doc.add_paragraph("[Primary Advisor/Internal Examiner]            [External Examiner]")
    doc.add_paragraph("[Official Title & Department]                 [Official Title & Department]")
    doc.add_paragraph("")
    doc.add_paragraph("[Chairman of the Department]                  [Dean of Faculty]")
    doc.add_paragraph("")
    doc.add_paragraph("Approval Date: [DATE]")
    doc.add_paragraph("")
    add_centered(doc, "DEPARTMENT OF [DEPARTMENT NAME]")
    add_centered(doc, "UNIVERSITY OF ENGINEERING & TECHNOLOGY, LAHORE")
    add_page_break(doc)

    # Copyright page
    add_centered(doc, "© [YEAR]")
    add_centered(doc, "[Scholar Full Name]")
    add_centered(doc, "All Rights Reserved.")
    add_centered(
        doc,
        "Any part of this thesis cannot be copied, reproduced, or published without the written approval of the Scholar.",
    )
    add_page_break(doc)

    add_section(
        doc,
        "ABSTRACT",
        [
            "This thesis presents a real-time adaptive framework for heart disease risk detection that combines static machine learning inference, online learning, and concept drift monitoring in a single deployable system. The core motivation is that static models may degrade over time due to changing clinical data distributions, creating reliability concerns in long-term use.",
            "The implemented framework includes a reproducible preprocessing pipeline, baseline model training, adaptive online learning loop, drift detector comparison, synthetic drift scenario benchmarking, explainability artifacts, and low-latency API serving. The baseline model is implemented using a scikit-learn pipeline (StandardScaler + LogisticRegression), while the adaptive module is implemented using River-based online logistic regression with ADWIN, DDM, and Page-Hinkley drift detectors.",
            "A FastAPI backend and Streamlit dashboard are integrated for real-time interaction and monitoring. Experiments on UCI-format heart disease data, including official UCI id=45 (303 records), show strong baseline holdout performance (accuracy around 0.8689, ROC-AUC around 0.9502) and stable repeated cross-validation performance (around 0.83 mean accuracy). Adaptive online tests verified prediction-boundary shifts and drift-event signaling under controlled phase changes.",
            "The thesis concludes that the proposed framework addresses practical limitations of static diagnosis systems by enabling adaptive updates, drift-aware monitoring, and reproducible real-time deployment workflows.",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "ACKNOWLEDGMENTS",
        [
            "I am deeply grateful to Almighty Allah for guidance and strength throughout this research journey.",
            "I sincerely thank my supervisor, [SUPERVISOR NAME], for valuable mentorship, constructive feedback, and continuous support during all phases of this thesis.",
            "I also acknowledge the faculty and staff of [DEPARTMENT NAME], UET Lahore, for providing an enabling academic environment. I am thankful to my colleagues and friends for technical discussions and review support, and especially to my family for their prayers, encouragement, and patience.",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "STATEMENT OF ORIGINALITY",
        [
            "It is stated that the research work presented in this thesis consists of my own ideas and research work. The contributions and ideas from others have been duly acknowledged and cited in the dissertation. This complete thesis is written by me.",
            "[YOUR FULL NAME]",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "NOMENCLATURE",
        [
            "x: feature vector",
            "y: true class label",
            "y-hat: predicted class label",
            "P(y=1|x): predicted probability of positive class",
            "AUC: Area Under the ROC Curve",
            "FNR: False Negative Rate",
            "ADWIN: Adaptive Windowing drift detector",
            "DDM: Drift Detection Method",
            "API: Application Programming Interface",
            "UCI: University of California Irvine Repository",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "CHAPTER 1: INTRODUCTION",
        [
            "Heart disease is one of the leading causes of global mortality. Early and reliable risk prediction is essential for preventive intervention and clinical decision support.",
            "Most machine learning systems used for disease prediction are static models trained once on historical data. Over time, real-world data distributions may change, causing model performance degradation. This phenomenon motivates adaptive learning and drift-aware monitoring.",
            "This research proposes and implements a real-time adaptive framework that integrates baseline static prediction, online model updates, drift detection, and deployment-ready interfaces through API and dashboard components.",
            "Key objectives include: (i) robust preprocessing and baseline modeling, (ii) online adaptation using River, (iii) detector comparison across ADWIN/DDM/Page-Hinkley, (iv) practical serving interfaces, and (v) reproducible experimentation and reporting.",
        ],
    )

    add_section(
        doc,
        "1.1 Problem Statement",
        [
            "Static heart disease models can become outdated when stream characteristics change. Without adaptation, prediction quality declines and trust in decision support decreases.",
            "A deployable framework is required that can continuously learn from feedback, monitor drift signals, and remain operational in low-latency real-time settings.",
        ],
    )

    add_section(
        doc,
        "1.2 Research Questions",
        [
            "Can an online adaptive model update its behavior under changing data/label patterns?",
            "Can drift events be observed and quantified in real time through integrated detector monitoring?",
            "Can the framework provide deployment-grade workflows with reproducible setup, testing, and reporting?",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "CHAPTER 2: LITERATURE REVIEW",
        [
            "Heart disease prediction literature reports strong performance for Logistic Regression, Random Forest, SVM, and deep models, often on UCI-derived datasets. However, many studies remain offline and do not address deployment-time drift.",
            "Concept drift occurs when feature-target relationships shift over time. Online learning can address this by updating models incrementally as feedback arrives.",
            "Popular drift detectors include ADWIN, DDM, and Page-Hinkley. Their sensitivity and false alarm behavior differ by stream dynamics, requiring comparative evaluation instead of single-detector assumptions.",
            "A key gap remains in end-to-end frameworks that jointly deliver static prediction, adaptive updates, drift monitoring, API serving, and dashboard-level observability. This thesis addresses that gap at a prototype research level.",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "CHAPTER 3: PROPOSED FRAMEWORK AND METHODOLOGY",
        [
            "The proposed framework is implemented as a modular architecture with data, model, serving, monitoring, and dashboard layers.",
            "Data preprocessing normalizes UCI-format aliases into canonical schema, validates required features, handles missing values, and creates deterministic processed datasets.",
            "The baseline model uses StandardScaler and LogisticRegression in a scikit-learn pipeline. Model artifact and serving configuration (artifact path, threshold, model version) are stored for deployment.",
            "Adaptive learning uses River online logistic regression. Each step follows prequential logic: predict first, compare with true label, update metrics/drift detector, then learn from that instance.",
            "Drift detector comparison is implemented for ADWIN, DDM, and Page-Hinkley, with summary metrics, per-step progress CSVs, and plot artifacts.",
            "FastAPI exposes static endpoints (/predict, /model-info) and adaptive endpoints (/adaptive/predict, /adaptive/learn, /adaptive/status). Streamlit dashboard provides both file-based analytics and live adaptive API interaction.",
        ],
    )

    add_section(
        doc,
        "3.1 Implementation Components",
        [
            "Data schema and preprocessing: src/heart_disease_rt/data/*",
            "Baseline training and benchmarking: src/heart_disease_rt/models/* and scripts/train_baseline.py",
            "Adaptive loop and detector comparison: scripts/run_adaptive_loop.py and models/adaptive.py",
            "Serving layer: api/main.py, serving/predictor.py, serving/adaptive_service.py",
            "Dashboard: dashboard/app.py with live adaptive API panel",
            "Reproducibility: Makefile + scripts/bootstrap_and_run_all.sh + automated test suite",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "CHAPTER 4: RESULTS AND DISCUSSION",
        [
            "Experiments were conducted on UCI-format heart disease data, including official UCI dataset id=45 (303 records).",
            "Baseline holdout run produced strong metrics: accuracy approximately 0.8689, precision 0.8125, recall 0.9286, F1 0.8667, and ROC-AUC 0.9502.",
            "Repeated stratified cross-validation provided stable mean accuracy around 0.83, supporting reproducibility and reducing single-split bias.",
            "Adaptive detector comparison generated per-detector summaries and progress plots for ADWIN, DDM, and Page-Hinkley under a common evaluation pipeline.",
            "Live adaptive API validation using two-phase label-shift stream confirmed adaptation: in phase A (label=1) risk increased toward positive prediction; in phase B (label=0) risk decreased and prediction switched to negative, while DDM reported a drift event.",
            "Latency profiling on 50 requests showed low typical response times (P50/P95 in low milliseconds range), with occasional outliers expected in local runtime conditions.",
            "Overall, results indicate the framework solves the practical static-model gap by enabling online updates, drift-aware monitoring, and operational deployment tooling.",
        ],
    )

    add_section(
        doc,
        "4.1 Key Quantitative Highlights",
        [
            "Baseline holdout accuracy: 0.8689",
            "Baseline ROC-AUC: 0.9502",
            "Repeated-CV mean accuracy (Logistic Regression): approximately 0.8313",
            "Adaptive online test: seen_samples reached 240 with drift_events = 1 in controlled phase-shift run",
            "Latest latency report (50 requests): mean approximately 53.06 ms; P50 approximately 13.15 ms; P95 approximately 15.15 ms",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "CHAPTER 5: CONCLUSION AND FUTURE WORK",
        [
            "This thesis delivers a practical, reproducible, and deployment-oriented adaptive framework for heart disease risk prediction.",
            "The implementation demonstrates that static and adaptive paradigms can be combined in one system with real-time APIs, drift visibility, and dashboard-level interaction.",
            "Adaptive behavior has been empirically validated through controlled online learning scenarios where prediction decisions changed in response to shifted supervision.",
            "Future work includes persistent adaptive state across API restarts, detector hyperparameter calibration, cost-sensitive optimization for false-negative reduction, and external multi-center validation.",
        ],
    )
    add_page_break(doc)

    add_section(
        doc,
        "BIBLIOGRAPHY (PLACEHOLDER FORMAT)",
        [
            "[1] UCI Machine Learning Repository, Heart Disease Dataset.",
            "[2] A. Bifet and R. Gavalda, Adaptive Windowing methods for drift detection.",
            "[3] J. Gama et al., Concept Drift Adaptation Surveys.",
            "[4] River Online Machine Learning Documentation.",
            "[5] Scikit-learn, FastAPI, and Streamlit official documentation.",
            "Note: Replace with department-recommended style (e.g., IEEE) using your actual cited papers.",
        ],
    )

    add_section(
        doc,
        "APPENDICES",
        [
            "Appendix A: API request/response examples (static and adaptive endpoints).",
            "Appendix B: Full reproducibility commands (bootstrap script, Makefile targets).",
            "Appendix C: Additional result artifacts (CV folds, drift logs, explainability outputs).",
            "Appendix D: Dashboard screenshots and deployment notes.",
        ],
    )

    add_section(
        doc,
        "VITA",
        [
            "[YOUR FULL NAME] completed [PREVIOUS DEGREE] in [YEAR] from [INSTITUTION].",
            "The scholar worked on machine learning for healthcare analytics and real-time adaptive systems.",
            "This MPhil thesis at UET Lahore focused on adaptive heart disease risk detection, concept drift handling, and deployable monitoring architecture.",
        ],
    )

    return doc


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate UET thesis implementation draft DOCX.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/UET_THESIS_IMPLEMENTATION_DRAFT.docx"),
        help="Output .docx path.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(args.output)
    print(f"Generated thesis draft DOCX: {args.output}")


if __name__ == "__main__":
    main()
