#!/usr/bin/env python3
"""Build a Word (.docx) document from repository plaintext documentation."""

from __future__ import annotations

from pathlib import Path

from docx import Document


def main() -> None:
    source_path = Path("docs/PROJECT_SETUP_AND_TECHNICAL_EXPLANATION.doc")
    output_path = Path("docs/PROJECT_SETUP_AND_TECHNICAL_EXPLANATION.docx")

    lines = source_path.read_text(encoding="utf-8").splitlines()

    document = Document()
    for line in lines:
        if line.startswith("============================================================"):
            continue
        if line.startswith("------------------------------------------------------------"):
            continue
        if line.strip() == "":
            document.add_paragraph("")
            continue

        # Basic heading heuristics for readability in Word.
        stripped = line.strip()
        if stripped[:2].isdigit() and ")" in stripped[:5]:
            document.add_heading(stripped, level=1)
            continue
        if stripped[:3].count(".") == 1 and stripped.startswith("6."):
            document.add_heading(stripped, level=2)
            continue

        document.add_paragraph(line)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    print(f"Generated Word document: {output_path}")


if __name__ == "__main__":
    main()
